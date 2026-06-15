from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT_ROOT / "docs" / "manager-demo-runbook.md"
OUTPUT = PROJECT_ROOT / "docs" / "manager-demo-runbook.docx"


def set_cell_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def add_code_block(document: Document, lines: list[str]):
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.right_indent = Inches(0.1)
        set_cell_shading(paragraph, "F3F4F6")
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(17, 24, 39)
    document.add_paragraph()


def add_markdown_line(document: Document, line: str):
    stripped = line.strip()
    if not stripped:
        return

    if stripped.startswith("# "):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(stripped[2:])
        run.font.name = "Arial"
        run.font.size = Pt(26)
        run.font.bold = False
        paragraph.paragraph_format.space_after = Pt(8)
        return

    if stripped.startswith("## "):
        paragraph = document.add_heading(stripped[3:], level=1)
        paragraph.paragraph_format.space_before = Pt(14)
        paragraph.paragraph_format.space_after = Pt(6)
        return

    if stripped.startswith("### "):
        paragraph = document.add_heading(stripped[4:], level=2)
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        return

    if stripped.startswith("- "):
        paragraph = document.add_paragraph(stripped[2:], style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        return

    if stripped.startswith("> "):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run(stripped[2:])
        run.italic = True
        run.font.color.rgb = RGBColor(75, 85, 99)
        return

    paragraph = document.add_paragraph(stripped)
    paragraph.paragraph_format.space_after = Pt(6)


def build_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 1"].font.color.rgb = RGBColor(17, 24, 39)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.color.rgb = RGBColor(31, 41, 55)

    code_lines = []
    in_code_block = False

    for line in SOURCE.read_text(encoding="utf-8").splitlines():
        if line.startswith("```"):
            if in_code_block:
                add_code_block(document, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        add_markdown_line(document, line)

    if code_lines:
        add_code_block(document, code_lines)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
