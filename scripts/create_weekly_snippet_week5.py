from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT_ROOT / "docs" / "weekly-snippet-through-week5.md"
DOCX_OUTPUT = PROJECT_ROOT / "docs" / "weekly-snippet-through-week5.docx"
PDF_OUTPUT = PROJECT_ROOT / "docs" / "weekly-snippet-through-week5.pdf"


def parse_markdown_lines():
    return SOURCE.read_text(encoding="utf-8").splitlines()


def add_docx_content(lines: list[str]):
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    pending_bullets = []

    def flush_bullets():
        nonlocal pending_bullets
        for bullet in pending_bullets:
            document.add_paragraph(bullet, style="List Bullet")
        pending_bullets = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_bullets()
            continue

        if stripped.startswith("# "):
            flush_bullets()
            document.add_heading(stripped[2:], level=0)
        elif stripped.startswith("## "):
            flush_bullets()
            document.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            flush_bullets()
            document.add_heading(stripped[4:], level=2)
        elif stripped.startswith("- "):
            pending_bullets.append(stripped[2:])
        else:
            flush_bullets()
            document.add_paragraph(stripped)

    flush_bullets()
    document.save(DOCX_OUTPUT)


def add_pdf_content(lines: list[str]):
    styles = getSampleStyleSheet()
    story = []
    pending_bullets = []

    def flush_bullets():
        nonlocal pending_bullets
        if pending_bullets:
            story.append(ListFlowable(
                [
                    ListItem(Paragraph(bullet, styles["BodyText"]))
                    for bullet in pending_bullets
                ],
                bulletType="bullet",
                leftIndent=18,
            ))
            story.append(Spacer(1, 8))
        pending_bullets = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush_bullets()
            continue

        if stripped.startswith("# "):
            flush_bullets()
            story.append(Paragraph(stripped[2:], styles["Title"]))
            story.append(Spacer(1, 8))
        elif stripped.startswith("## "):
            flush_bullets()
            story.append(Paragraph(stripped[3:], styles["Heading1"]))
            story.append(Spacer(1, 6))
        elif stripped.startswith("### "):
            flush_bullets()
            story.append(Paragraph(stripped[4:], styles["Heading2"]))
            story.append(Spacer(1, 4))
        elif stripped.startswith("- "):
            pending_bullets.append(stripped[2:])
        else:
            flush_bullets()
            story.append(Paragraph(stripped, styles["BodyText"]))
            story.append(Spacer(1, 5))

    flush_bullets()
    SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=letter,
        rightMargin=42,
        leftMargin=42,
        topMargin=42,
        bottomMargin=42,
    ).build(story)


def main():
    lines = parse_markdown_lines()
    add_docx_content(lines)
    add_pdf_content(lines)
    print(DOCX_OUTPUT)
    print(PDF_OUTPUT)


if __name__ == "__main__":
    main()
